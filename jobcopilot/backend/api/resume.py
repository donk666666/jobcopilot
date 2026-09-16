"""
简历优化 API
"""

import json
import re
from fastapi import APIRouter, HTTPException, Depends
from pydantic import BaseModel
from sqlalchemy.orm import Session

from agent.core import match_resume_direct, tailor_resume_direct
import config
from config import DEEPSEEK_API_KEY, DEEPSEEK_BASE_URL, DEEPSEEK_MODEL
from rag.vector_store import get_vector_store
from database import get_db, ResumeOptimization, Resume, User
from langchain_openai import ChatOpenAI
from langchain_core.messages import HumanMessage, SystemMessage
from prompts.resume_tailor import RESUME_MATCH_SYSTEM, RESUME_MATCH_USER_TEMPLATE, RESUME_TAILOR_SYSTEM, RESUME_TAILOR_USER_TEMPLATE
from cache import cache_llm_result, get_cached_llm, get_redis_status
from api.auth import get_current_user

router = APIRouter(prefix="/api/resume", tags=["简历优化"])


class ResumeMatchRequest(BaseModel):
    resume_text: str
    jd_analysis: str
    jd_analysis_id: int | None = None


class ResumeTailorRequest(BaseModel):
    resume_text: str
    jd_analysis: str
    match_result: str
    opt_id: int | None = None
    show_annotations: bool = True


@router.post("/match")
async def match_resume(request: ResumeMatchRequest, db: Session = Depends(get_db)):
    """评估简历与职位的匹配度，结果自动入库"""
    if not request.resume_text.strip():
        raise HTTPException(status_code=400, detail="简历文本不能为空")

    # 查缓存
    cached = get_cached_llm(request.jd_analysis, request.resume_text, "match")
    if cached is not None:
        parsed = cached
        record = ResumeOptimization(
            jd_analysis_id=request.jd_analysis_id,
            resume_text=request.resume_text,
            jd_analysis_json=request.jd_analysis,
            match_score=parsed.get("match_score"),
            match_detail=json.dumps(parsed, ensure_ascii=False),
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        return {"result": parsed, "success": True, "opt_id": record.id, "cached": True}

    result = await match_resume_direct(
        jd_analysis=request.jd_analysis,
        resume_text=request.resume_text,
        api_key=DEEPSEEK_API_KEY,
        base_url=DEEPSEEK_BASE_URL,
    )

    if result.get("success") and result.get("result"):
        parsed = result["result"]
        cache_llm_result(request.jd_analysis, request.resume_text, "match", parsed)
        record = ResumeOptimization(
            jd_analysis_id=request.jd_analysis_id,
            resume_text=request.resume_text,
            jd_analysis_json=request.jd_analysis,
            match_score=parsed.get("match_score"),
            match_detail=json.dumps(parsed, ensure_ascii=False),
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        return {"result": parsed, "success": True, "opt_id": record.id}

    return result


@router.post("/tailor")
async def tailor_resume(request: ResumeTailorRequest, db: Session = Depends(get_db)):
    """根据JD定向优化简历，结果自动入库"""
    if not request.resume_text.strip():
        raise HTTPException(status_code=400, detail="简历文本不能为空")

    # 使用自定义模板以支持注释开关
    annotations_on = request.show_annotations

    tailor_system = RESUME_TAILOR_SYSTEM
    if not annotations_on:
        tailor_system = RESUME_TAILOR_SYSTEM.replace(
            "并在每处修改之后用【改动说明：xxx】标注", "。不要添加任何改动说明或注释"
        )

    # 简历过长截断（避免 LLM 输出耗时）—— 一般简历 2500 字足够保留所有关键信息
    resume_text = request.resume_text
    if len(resume_text) > 2500:
        resume_text = resume_text[:2500] + "\n...(简历后段省略)"

    # 查缓存
    cached = get_cached_llm(request.jd_analysis, resume_text, f"tailor_ann{1 if annotations_on else 0}")
    if cached is not None:
        # 存入库
        if request.opt_id:
            record = db.query(ResumeOptimization).filter(ResumeOptimization.id == request.opt_id).first()
            if record:
                record.tailored_resume = cached
                record.annotations_enabled = 1 if annotations_on else 0
                db.commit()
                return {"result": cached, "success": True, "opt_id": record.id, "cached": True}

        # 无 opt_id，新建
        record = ResumeOptimization(
            jd_analysis_json=request.jd_analysis,
            resume_text=resume_text,
            tailored_resume=cached,
            annotations_enabled=1 if annotations_on else 0,
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        return {"result": cached, "success": True, "opt_id": record.id, "cached": True}

    llm = ChatOpenAI(
        api_key=DEEPSEEK_API_KEY,
        base_url=DEEPSEEK_BASE_URL,
        model=DEEPSEEK_MODEL,
        temperature=0.6,
        max_tokens=2500,  # 限制输出长度，避免 LLM 啰嗦
    )

    messages = [
        SystemMessage(content=tailor_system),
        HumanMessage(content=RESUME_TAILOR_USER_TEMPLATE.format(
            jd_analysis=request.jd_analysis,
            match_result=request.match_result,
            resume_text=resume_text
        ))
    ]

    response = await llm.ainvoke(messages)
    tailored = response.content

    cache_llm_result(request.jd_analysis, resume_text, f"tailor_ann{1 if annotations_on else 0}", tailored)

    # 存入库
    if request.opt_id:
        record = db.query(ResumeOptimization).filter(ResumeOptimization.id == request.opt_id).first()
        if record:
            record.tailored_resume = tailored
            record.annotations_enabled = 1 if annotations_on else 0
            db.commit()
            return {"result": tailored, "success": True, "opt_id": record.id}

    # 无 opt_id，新建
    record = ResumeOptimization(
        jd_analysis_json=request.jd_analysis,
        resume_text=resume_text,
        tailored_resume=tailored,
        annotations_enabled=1 if annotations_on else 0,
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return {"result": tailored, "success": True, "opt_id": record.id}


class FullPipelineRequest(BaseModel):
    resume_text: str
    jd_text: str
    jd_analysis_id: int | None = None
    style: str = "professional"
    candidate_name: str = ""


@router.post("/full-pipeline")
async def start_full_pipeline(request: FullPipelineRequest):
    """一键全流程：JD分析 → 匹配 → 定制简历 → 求职信（Celery 异步）"""
    try:
        from tasks import run_full_pipeline
    except ImportError:
        raise HTTPException(
            status_code=503,
            detail="异步任务服务未启动，请使用分步操作（JD分析 → 简历匹配 → 简历定制）",
        )

    try:
        import redis as redis_lib
        # 用 RESP2 协议（Redis 5.x 不支持 HELLO 3）
        r = redis_lib.Redis.from_url(
            config.REDIS_URL + "/1" if not config.REDIS_URL.endswith("/1") else config.REDIS_URL,
            socket_connect_timeout=2,
            socket_timeout=2,
            protocol=2,
        )
        r.ping()
    except Exception:
        raise HTTPException(
            status_code=503,
            detail="消息队列服务不可用，请使用分步操作",
        )

    task = run_full_pipeline.delay(
        resume_text=request.resume_text,
        jd_text=request.jd_text,
        jd_analysis_id=request.jd_analysis_id,
        style=request.style,
        candidate_name=request.candidate_name,
    )

    return {"task_id": task.id, "status": "pending"}


@router.get("/redis-status")
def redis_status():
    """查询 Redis 缓存连接状态"""
    return get_redis_status()


@router.get("/history")
def list_optimizations(db: Session = Depends(get_db)):
    """获取简历优化历史"""
    records = db.query(ResumeOptimization).order_by(ResumeOptimization.updated_at.desc()).limit(20).all()
    return [
        {
            "id": r.id,
            "jd_analysis_id": r.jd_analysis_id,
            "match_score": r.match_score,
            "match_detail": json.loads(r.match_detail) if r.match_detail else None,
            "tailored_resume": r.tailored_resume,
            "annotations_enabled": bool(r.annotations_enabled),
            "resume_preview": r.resume_text[:80] + ("..." if len(r.resume_text) > 80 else ""),
            "created_at": r.created_at.isoformat() if r.created_at else "",
            "updated_at": r.updated_at.isoformat() if r.updated_at else "",
        }
        for r in records
    ]


@router.get("/history/{opt_id}")
def get_optimization(opt_id: int, db: Session = Depends(get_db)):
    """获取单条优化记录"""
    r = db.query(ResumeOptimization).filter(ResumeOptimization.id == opt_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="记录不存在")
    return {
        "id": r.id,
        "jd_analysis_id": r.jd_analysis_id,
        "resume_text": r.resume_text,
        "jd_analysis_json": r.jd_analysis_json,
        "match_score": r.match_score,
        "match_detail": json.loads(r.match_detail) if r.match_detail else None,
        "tailored_resume": r.tailored_resume,
        "annotations_enabled": bool(r.annotations_enabled),
        "created_at": r.created_at.isoformat() if r.created_at else "",
    }


@router.delete("/history/{opt_id}")
def delete_optimization(opt_id: int, db: Session = Depends(get_db)):
    """删除一条优化记录"""
    r = db.query(ResumeOptimization).filter(ResumeOptimization.id == opt_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="记录不存在")
    db.delete(r)
    db.commit()
    return {"message": "已删除"}


class ResumeSaveRequest(BaseModel):
    resume_text: str
    name: str = "默认简历"


@router.get("/active")
def get_active_resume(db: Session = Depends(get_db), current: User = Depends(get_current_user)):
    """获取当前用户的激活简历"""
    r = db.query(Resume).filter(
        Resume.user_id == current.id,
        Resume.is_active == 1,
    ).order_by(Resume.updated_at.desc()).first()
    if not r:
        return {"content": "", "name": ""}
    return {"id": r.id, "content": r.content, "name": r.name}


@router.post("/active")
def save_active_resume(req: ResumeSaveRequest, db: Session = Depends(get_db), current: User = Depends(get_current_user)):
    """保存当前用户的激活简历（同时取消该用户其他简历的激活）"""
    # 先取消该用户所有激活
    db.query(Resume).filter(
        Resume.user_id == current.id,
        Resume.is_active == 1,
    ).update({"is_active": 0})
    # 写入新的激活简历
    new_r = Resume(user_id=current.id, name=req.name, content=req.resume_text, is_active=1)
    db.add(new_r)
    db.commit()
    db.refresh(new_r)
    return {"id": new_r.id, "name": new_r.name, "message": "已保存为激活简历"}


@router.get("/versions")
def list_resumes(db: Session = Depends(get_db)):
    """获取所有历史简历（按更新时间倒序）"""
    records = db.query(Resume).order_by(Resume.updated_at.desc()).all()
    return [
        {
            "id": r.id,
            "name": r.name,
            "content": r.content,
            "is_active": r.is_active,
            "updated_at": r.updated_at.isoformat() if r.updated_at else "",
        }
        for r in records
    ]


@router.put("/{resume_id}")
def update_resume(resume_id: int, request: ResumeSaveRequest, db: Session = Depends(get_db)):
    """修改历史简历文本（按 id 更新内容/名称）"""
    r = db.query(Resume).filter(Resume.id == resume_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="简历不存在")
    r.content = request.resume_text
    if request.name:
        r.name = request.name
    db.commit()
    db.refresh(r)
    return {"id": r.id, "message": "已更新"}


@router.delete("/{resume_id}")
def delete_resume(resume_id: int, db: Session = Depends(get_db)):
    """删除历史简历"""
    r = db.query(Resume).filter(Resume.id == resume_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="简历不存在")
    db.delete(r)
    db.commit()
    return {"id": resume_id, "message": "已删除"}


@router.post("/active")
def save_active_resume(request: ResumeSaveRequest, db: Session = Depends(get_db)):
    """保存简历（设为激活，其他取消激活）"""
    # 取消所有激活
    db.query(Resume).filter(Resume.is_active == 1).update({Resume.is_active: 0})
    # 新建一条激活记录
    record = Resume(name=request.name, content=request.resume_text, is_active=1)
    db.add(record)
    db.commit()
    db.refresh(record)
    return {"id": record.id, "message": "已保存"}


@router.get("/rag-stats")
async def get_rag_stats():
    """获取RAG向量库状态"""
    store = get_vector_store()
    return store.get_store_stats()


# ===== 简历导出 Word =====

class ExportDocxRequest(BaseModel):
    title: str = "简历"
    content: str
    candidate_name: str | None = None


@router.post("/export-docx")
async def export_resume_docx(req: ExportDocxRequest):
    """把改写后的简历文本导出为 Word .docx 文件"""
    from io import BytesIO
    from fastapi.responses import StreamingResponse
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # 设置默认字体（中文兼容：同时设置 ASCII + 东亚字体回退）
    from docx.oxml.ns import qn
    style = doc.styles['Normal']
    style.font.name = 'Calibri'  # ASCII 默认
    style.font.size = Pt(11)
    # 东亚字体回退（中文）
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), '宋体')  # 中文回退字体
    rfonts.set(qn('w:ascii'), 'Calibri')
    rfonts.set(qn('w:hAnsi'), 'Calibri')

    content = (req.content or "").strip()
    if not content:
        content = "（暂无内容）"

    title = req.title or "个人简历"
    # 主标题
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = head.add_run(title)
    run.bold = True
    run.font.size = Pt(18)

    # 候选姓名
    if req.candidate_name:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub.add_run(req.candidate_name)
        sub_run.font.size = Pt(12)
        sub_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph()  # 空行

    # 按行解析内容（Markdown 简易转换）
    lines = content.split("\n")
    for line in lines:
        line = line.rstrip()
        if not line.strip():
            doc.add_paragraph()
            continue
        # 去除 Markdown 标记
        text = line.strip()
        is_heading = False
        if text.startswith("### "):
            text = text[4:]; is_heading = True
        elif text.startswith("## "):
            text = text[3:]; is_heading = True
        elif text.startswith("# "):
            text = text[2:]; is_heading = True
        # 处理加粗 **xxx**
        import re
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        # 列表 - xxx
        m = re.match(r'^[-*]\s+(.+)$', text)
        if m:
            text = "• " + m.group(1)

        p = doc.add_paragraph()
        if is_heading:
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x00, 0xb5, 0x74)
        else:
            p.add_run(text)

    # 输出
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    base_name = req.candidate_name or "resume"
    ascii_name = "".join(c if ord(c) < 128 else "_" for c in base_name)
    utf8_name = f"{base_name}_简历.docx"
    import urllib.parse
    encoded_utf8 = urllib.parse.quote(utf8_name)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{ascii_name}_resume.docx"; '
                f"filename*=UTF-8''{encoded_utf8}"
            ),
        },
    )
